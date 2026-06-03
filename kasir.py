import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import os, sys
from datetime import datetime
from collections import defaultdict

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Install openpyxl: pip install openpyxl")
    sys.exit(1)
try:
    import pandas as pd
except ImportError:
    print("Install pandas: pip install pandas")
    sys.exit(1)
try:
    import numpy as np
except ImportError:
    print("Install numpy: pip install numpy")
    sys.exit(1)
try:
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
    import matplotlib.ticker as mticker
except ImportError:
    print("Install matplotlib: pip install matplotlib")
    sys.exit(1)

# ═════════════════════════════════════════════════════════════════════════════
#  OOP: DATA CLASSES
# ═════════════════════════════════════════════════════════════════════════════
class Barang:
    """Merepresentasikan satu jenis barang di toko."""
    def __init__(self, kode: str, nama: str, harga: int, kategori: str = "lainnya"):
        self.kode      = kode
        self.nama      = nama
        self.harga     = harga
        self.kategori  = kategori

    def __repr__(self):
        return f"Barang({self.kode}, {self.nama}, Rp{self.harga:,})"


class ItemTransaksi:
    """Satu baris item dalam transaksi (barang + qty)."""
    def __init__(self, barang: Barang, qty: int = 1):
        self.barang = barang
        self.qty    = qty

    @property
    def subtotal(self) -> int:
        return self.barang.harga * self.qty


class Transaksi:
    """Satu sesi transaksi kasir."""
    _counter = 0

    def __init__(self, tanggal: str, metode: str):
        Transaksi._counter += 1
        self.no_transaksi = f"TRX/{tanggal.replace('-','')}/{str(Transaksi._counter).zfill(3)}"
        self.tanggal      = tanggal
        self.waktu        = datetime.now().strftime("%H:%M:%S")
        self.metode       = metode
        self.items: list[ItemTransaksi] = []

    def tambah_item(self, barang: Barang, qty: int = 1):
        for it in self.items:
            if it.barang.kode == barang.kode:
                it.qty += qty
                return
        self.items.append(ItemTransaksi(barang, qty))

    def hapus_item(self, kode: str):
        self.items = [it for it in self.items if it.barang.kode != kode]

    @property
    def total(self) -> int:
        return sum(it.subtotal for it in self.items)

    def buat_struk(self, uang_bayar: int = 0) -> str:
        garis  = "=" * 38
        garisT = "-" * 38
        baris  = [garis,
                  "         JACK(OWI) MART",
                  "    Jl. Raya No.1, Jakarta Pusat",
                  "      Telp: (021) 1234-5678",
                  garisT,
                  f"  No  : {self.no_transaksi}",
                  f"  Tgl : {self.tanggal}   {self.waktu}",
                  f"  Bayar: {self.metode.upper()}",
                  garis,
                  f"{'Barang':<20}{'Qty':>4}{'Harga':>13}",
                  garisT]
        for it in self.items:
            nm = it.barang.nama[:19]
            baris.append(f"{nm:<20}{it.qty:>4}{self._rp(it.subtotal):>13}")
            baris.append(f"  @ {self._rp(it.barang.harga)} x {it.qty}")
        baris += [garis,
                  f"{'TOTAL':<25}{self._rp(self.total):>13}",
                  garisT]
        if self.metode.lower() == "cash" and uang_bayar > 0:
            baris.append(f"  Tunai   : {self._rp(uang_bayar)}")
            baris.append(f"  Kembali : {self._rp(uang_bayar - self.total)}")
        else:
            baris.append("  *** PEMBAYARAN QRIS BERHASIL ***")
        baris += [garis,
                  "    Terima Kasih Sudah Belanja!",
                  "     Selamat Datang Kembali :)",
                  garis]
        return "\n".join(baris)

    @staticmethod
    def _rp(n: int) -> str:
        return f"Rp{n:,}".replace(",", ".")


# ═════════════════════════════════════════════════════════════════════════════
#  FILE HANDLING (pandas + openpyxl)
# ═════════════════════════════════════════════════════════════════════════════
class FileHandler:
    """Baca/tulis file Excel untuk items dan penjualan."""

    # ── Temukan folder program ────────────────────────────────────────────────
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

    @classmethod
    def path_items(cls) -> str:
        return os.path.join(cls.BASE_DIR, "items_barang.xlsx")

    @classmethod
    def path_penjualan(cls, tanggal: str) -> str:
        bulan = tanggal[:7]          # YYYY-MM
        fname = f"penjualan_{bulan}.xlsx"
        return os.path.join(cls.BASE_DIR, fname)

    # ── Baca daftar barang ────────────────────────────────────────────────────
    @classmethod
    def baca_barang(cls) -> list[Barang]:
        path = cls.path_items()
        if not os.path.exists(path):
            raise FileNotFoundError(
                f"File '{path}' tidak ditemukan!\n"
                "Jalankan 'python buat_excel.py' terlebih dahulu."
            )
        df = pd.read_excel(path, sheet_name="Daftar Barang",
                           skiprows=1, dtype={"Kode": str})
        df.columns = df.columns.str.strip()
        kolom_map = {
            "Kode": "kode", "Nama Barang": "nama",
            "Harga (Rp)": "harga", "Kategori": "kategori",
        }
        df = df.rename(columns=kolom_map)
        df = df.dropna(subset=["kode", "nama", "harga"])
        df["harga"]    = pd.to_numeric(df["harga"], errors="coerce").fillna(0).astype(int)
        df["kategori"] = df.get("kategori", pd.Series(["lainnya"]*len(df))).fillna("lainnya")
        df["kode"]     = df["kode"].astype(str).str.strip()
        return [Barang(r.kode, r.nama, r.harga, r.kategori)
                for r in df.itertuples(index=False)
                if r.kode and str(r.kode).lower() != "total"]

    # ── Simpan transaksi ke Excel ─────────────────────────────────────────────
    @classmethod
    def simpan_transaksi(cls, trx: Transaksi):
        path   = cls.path_penjualan(trx.tanggal)
        biru   = "0D47A1"; biru_m = "1565C0"; biru_l = "BBDEFB"
        kuning = "FFD54F"; putih  = "FFFFFF"; abu    = "F5F5F5"

        def border():
            s = Side(style='thin', color='BDBDBD')
            return Border(left=s, right=s, top=s, bottom=s)

        def style_header(c, bold=False, bg=biru_m, fg=putih):
            c.font      = Font(name="Times New Roman", size=10, bold=bold, color=fg)
            c.fill      = PatternFill("solid", fgColor=bg)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border    = border()

        # ── Buka / buat workbook ──────────────────────────────────────────────
        if os.path.exists(path):
            wb = openpyxl.load_workbook(path)
        else:
            wb = openpyxl.Workbook()
            for sname in list(wb.sheetnames):
                del wb[sname]

        # ── Sheet: Log Transaksi ──────────────────────────────────────────────
        if "Log Transaksi" not in wb.sheetnames:
            ws1 = wb.create_sheet("Log Transaksi")
            ws1.merge_cells("A1:I1")
            ws1["A1"] = "JACK(OWI) MART - LOG TRANSAKSI PENJUALAN"
            ws1["A1"].font      = Font(name="Times New Roman", size=13, bold=True, color=kuning)
            ws1["A1"].fill      = PatternFill("solid", fgColor=biru)
            ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws1.row_dimensions[1].height = 28

            hdrs = ["No","Tanggal","No Transaksi","Waktu","Metode","Kode","Nama Barang","Qty","Subtotal (Rp)"]
            for col, h in enumerate(hdrs, 1):
                style_header(ws1.cell(row=2, column=col, value=h))
            ws1.row_dimensions[2].height = 24
            ws1.freeze_panes = "A3"
            for w, col in zip([6,14,28,10,10,8,36,6,18], range(1,10)):
                ws1.column_dimensions[get_column_letter(col)].width = w
        else:
            ws1 = wb["Log Transaksi"]

        # Tentukan baris mulai
        start = ws1.max_row + 1 if ws1.max_row > 2 else 3
        no_awal = ws1.max_row - 2 if ws1.max_row > 2 else 0

        for i, it in enumerate(trx.items):
            row = start + i
            no  = no_awal + i + 1
            bg  = abu if no % 2 == 0 else putih
            vals = [no, trx.tanggal, trx.no_transaksi, trx.waktu, trx.metode,
                    it.barang.kode, it.barang.nama, it.qty, it.subtotal]
            for col, val in enumerate(vals, 1):
                c = ws1.cell(row=row, column=col, value=val)
                c.font   = Font(name="Times New Roman", size=10)
                c.fill   = PatternFill("solid", fgColor=bg)
                c.border = border()
                if col in [1, 4, 5, 6, 8]:
                    c.alignment = Alignment(horizontal="center")
                elif col == 9:
                    c.alignment   = Alignment(horizontal="right")
                    c.number_format = '#,##0'

        # ── Sheet: Rekap Harian ───────────────────────────────────────────────
        tgl = trx.tanggal
        if "Rekap Harian" not in wb.sheetnames:
            ws2 = wb.create_sheet("Rekap Harian")
            ws2.merge_cells("A1:E1")
            ws2["A1"] = "REKAP PENJUALAN HARIAN"
            ws2["A1"].font      = Font(name="Times New Roman", size=13, bold=True, color=kuning)
            ws2["A1"].fill      = PatternFill("solid", fgColor=biru)
            ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws2.row_dimensions[1].height = 28
            h2 = ["Tanggal","Jml Transaksi","Total Item","Total Pendapatan (Rp)","Update Terakhir"]
            for col, h in enumerate(h2, 1):
                style_header(ws2.cell(row=2, column=col, value=h))
            ws2.row_dimensions[2].height = 24
            ws2.freeze_panes = "A3"
            for w, col in zip([16,16,14,24,20], range(1,6)):
                ws2.column_dimensions[get_column_letter(col)].width = w
        else:
            ws2 = wb["Rekap Harian"]

        # Cari / buat baris untuk tanggal ini
        tgl_row = None
        for row in ws2.iter_rows(min_row=3, max_col=1):
            if str(row[0].value) == str(tgl):
                tgl_row = row[0].row
                break
        if tgl_row is None:
            tgl_row = ws2.max_row + 1 if ws2.max_row >= 3 else 3
            ws2.cell(row=tgl_row, column=1, value=tgl)

        # Hitung dari log
        ws1_data = [[ws1.cell(row=r, column=c).value
                     for c in range(1, 10)]
                    for r in range(3, ws1.max_row + 1)]
        baris_tgl = [r for r in ws1_data if str(r[1]) == str(tgl) and r[2]]
        trx_unik  = len(set(r[2] for r in baris_tgl))
        total_qty = sum(int(r[7]) for r in baris_tgl if r[7])
        total_rp  = sum(int(r[8]) for r in baris_tgl if r[8])

        bg2 = abu if (tgl_row - 2) % 2 == 0 else putih
        for col, val in enumerate([tgl, trx_unik, total_qty, total_rp,
                                    datetime.now().strftime("%Y-%m-%d %H:%M")], 1):
            c = ws2.cell(row=tgl_row, column=col, value=val)
            c.font   = Font(name="Times New Roman", size=10)
            c.fill   = PatternFill("solid", fgColor=bg2)
            c.border = border()
            if col == 4:
                c.alignment   = Alignment(horizontal="right")
                c.number_format = '#,##0'
            else:
                c.alignment = Alignment(horizontal="center")

        wb.save(path)
        return path


# ═════════════════════════════════════════════════════════════════════════════
#  GRAFIK (Matplotlib + NumPy)
# ═════════════════════════════════════════════════════════════════════════════
class GrafikPenjualan:
    """Tampilkan grafik penjualan dari file Excel bulan ini."""

    @staticmethod
    def tampilkan(parent_frame, tanggal: str):
        path = FileHandler.path_penjualan(tanggal)
        for w in parent_frame.winfo_children():
            w.destroy()

        if not os.path.exists(path):
            tk.Label(parent_frame, text="Belum ada data penjualan bulan ini.",
                     font=("Times New Roman", 11), fg="#1565C0").pack(pady=20)
            return

        try:
            df = pd.read_excel(path, sheet_name="Rekap Harian", skiprows=1)
            df.columns = df.columns.str.strip()
            df = df.dropna(subset=[df.columns[0]])
            df = df[df[df.columns[0]].astype(str).str.match(r'\d{4}-\d{2}-\d{2}')]

            if df.empty:
                tk.Label(parent_frame, text="Data rekap kosong.",
                         font=("Times New Roman", 11)).pack(pady=20)
                return

            tgl_col  = df.columns[0]
            pend_col = [c for c in df.columns if "Pendapatan" in str(c) or "Total" in str(c)][-1]
            trx_col  = [c for c in df.columns if "Transaksi" in str(c)][0]

            labels   = df[tgl_col].astype(str).tolist()
            pend     = pd.to_numeric(df[pend_col], errors="coerce").fillna(0).values
            trx_cnt  = pd.to_numeric(df[trx_col],  errors="coerce").fillna(0).values

            # NumPy: statistik
            rata2 = np.mean(pend)
            maks  = np.max(pend)
            total = np.sum(pend)

            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(9, 3.5), dpi=90)
            fig.patch.set_facecolor("#E3F2FD")

            # Bar chart pendapatan
            colors = ["#1565C0" if p < rata2 else "#0D47A1" for p in pend]
            bars = ax1.bar(range(len(labels)), pend / 1000, color=colors,
                           edgecolor="#FFD54F", linewidth=0.8)
            ax1.axhline(rata2 / 1000, color="#FFD54F", linestyle="--",
                        linewidth=1.2, label=f"Rata-rata: Rp{rata2:,.0f}")
            ax1.set_xticks(range(len(labels)))
            ax1.set_xticklabels([l[-5:] for l in labels], rotation=45,
                                 fontsize=7, fontname="Times New Roman")
            ax1.yaxis.set_major_formatter(mticker.FuncFormatter(
                lambda x, _: f"Rp{x:.0f}K"))
            ax1.set_title("Pendapatan Harian", fontname="Times New Roman",
                          fontsize=10, color="#0D47A1", fontweight="bold")
            ax1.set_facecolor("#F0F8FF")
            ax1.legend(fontsize=7)
            ax1.tick_params(labelsize=7)

            # Pie chart transaksi per hari
            if np.sum(trx_cnt) > 0:
                wedge_colors = plt.cm.Blues(np.linspace(0.4, 0.9, len(labels)))
                ax2.pie(trx_cnt, labels=[l[-5:] for l in labels],
                        autopct="%1.0f%%", colors=wedge_colors,
                        textprops={"fontsize": 7, "fontname": "Times New Roman"},
                        startangle=90)
                ax2.set_title("Distribusi Transaksi", fontname="Times New Roman",
                              fontsize=10, color="#0D47A1", fontweight="bold")
            else:
                ax2.text(0.5, 0.5, "Tidak ada data", ha="center", va="center",
                         fontname="Times New Roman")

            fig.suptitle(
                f"Total Bulan: Rp{total:,.0f}  |  Maks/Hari: Rp{maks:,.0f}",
                fontname="Times New Roman", fontsize=9, color="#0D47A1", y=1.01
            )
            plt.tight_layout()

            canvas = FigureCanvasTkAgg(fig, master=parent_frame)
            canvas.draw()
            canvas.get_tk_widget().pack(fill="both", expand=True)

        except Exception as e:
            tk.Label(parent_frame, text=f"Gagal memuat grafik:\n{e}",
                     font=("Times New Roman", 10), fg="red").pack(pady=10)


# ═════════════════════════════════════════════════════════════════════════════
#  GUI UTAMA
# ═════════════════════════════════════════════════════════════════════════════
C_BG      = "#E3F2FD"
C_HEADER  = "#0D47A1"
C_MID     = "#1565C0"
C_LIGHT   = "#BBDEFB"
C_KUNING  = "#FFD54F"
C_PUTIH   = "#FAFAFA"
C_TEXT    = "#0D47A1"
FONT_J    = ("Times New Roman", 10)
FONT_B    = ("Times New Roman", 10, "bold")
FONT_BIG  = ("Times New Roman", 13, "bold")

class KasirApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Jack(owi) Mart - Sistem Kasir Digital")
        self.configure(bg=C_BG)
        self.state("zoomed")          # Full screen

        # State
        self._barang_list: list[Barang] = []
        self._filtered:    list[Barang] = []
        self._transaksi:   Transaksi | None = None
        self._uang_bayar   = 0
        self._trx_count    = 0
        self._trx_total    = 0
        self._kat_aktif    = "semua"

        self._build_ui()
        self._set_tanggal_hari_ini()
        self._muat_barang()

    # ── Build UI ──────────────────────────────────────────────────────────────
    def _build_ui(self):
        self._build_headline()
        self._build_datebar()

        body = tk.Frame(self, bg=C_BG)
        body.pack(fill="both", expand=True, padx=14, pady=10)
        body.columnconfigure(0, weight=2)
        body.columnconfigure(1, weight=1)
        body.rowconfigure(0, weight=1)

        self._build_left(body)
        self._build_right(body)

    def _build_headline(self):
        f = tk.Frame(self, bg=C_HEADER, pady=12)
        f.pack(fill="x")
        lbl = tk.Label(f, bg=C_HEADER)
        lbl.pack()
        tk.Label(lbl, text="Jack(owi)", bg=C_HEADER, fg=C_KUNING,
                 font=("Times New Roman", 26, "bold")).pack(side="left")
        tk.Label(lbl, text=" Mart", bg=C_HEADER, fg="#F5F5DC",
                 font=("Times New Roman", 26, "bold")).pack(side="left")
        tk.Label(f, text="SISTEM KASIR DIGITAL", bg=C_HEADER, fg="#90CAF9",
                 font=("Times New Roman", 9)).pack()

    def _build_datebar(self):
        f = tk.Frame(self, bg=C_MID, pady=6)
        f.pack(fill="x")
        tk.Label(f, text="📅 Tanggal:", bg=C_MID, fg="#90CAF9",
                 font=FONT_J).pack(side="left", padx=(16, 4))
        self.var_tgl = tk.StringVar()
        tk.Entry(f, textvariable=self.var_tgl, width=12,
                 font=FONT_J, bg=C_HEADER, fg=C_KUNING,
                 insertbackground=C_KUNING, relief="flat",
                 bd=4).pack(side="left")
        tk.Label(f, text="(YYYY-MM-DD)", bg=C_MID, fg="#64B5F6",
                 font=("Times New Roman", 8)).pack(side="left", padx=4)

        # Info transaksi
        self.lbl_trx_count = tk.Label(f, text="Transaksi: 0  |  Total: Rp 0",
                                       bg=C_MID, fg=C_KUNING, font=FONT_B)
        self.lbl_trx_count.pack(side="right", padx=16)

    def _frame_box(self, parent, title, row, col, rowspan=1, colspan=1,
                   sticky="nsew", pady_inner=8):
        outer = tk.Frame(parent, bg=C_LIGHT, bd=1, relief="flat")
        outer.grid(row=row, column=col, rowspan=rowspan, columnspan=colspan,
                   sticky=sticky, padx=6, pady=6)
        tk.Label(outer, text=title, bg=C_MID, fg=C_PUTIH,
                 font=FONT_B, pady=5).pack(fill="x")
        inner = tk.Frame(outer, bg=C_PUTIH, padx=8, pady=pady_inner)
        inner.pack(fill="both", expand=True)
        return inner

    # ── LEFT PANEL ────────────────────────────────────────────────────────────
    def _build_left(self, parent):
        left = tk.Frame(parent, bg=C_BG)
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
        left.rowconfigure(1, weight=2)
        left.columnconfigure(0, weight=1)

        # Cari barang
        f_cari = self._frame_box(left, "🔍  Cari & Input Barang", 0, 0)
        self._build_cari(f_cari)

        # Keranjang
        f_cart = self._frame_box(left, "🛒  Keranjang Belanja", 1, 0)
        self._build_cart(f_cart)

        # Action bar
        f_act = tk.Frame(left, bg=C_LIGHT, bd=1, relief="flat")
        f_act.grid(row=2, column=0, sticky="ew", padx=6, pady=6)
        self._build_action(f_act)

    def _build_cari(self, parent):
        row1 = tk.Frame(parent, bg=C_PUTIH)
        row1.pack(fill="x", pady=(0, 6))

        self.var_kode = tk.StringVar()
        e = tk.Entry(row1, textvariable=self.var_kode, font=FONT_J,
                     width=28, bg="#F0F8FF", fg=C_TEXT,
                     relief="solid", bd=1)
        e.pack(side="left", padx=(0, 6), ipady=4)
        e.bind("<Return>", lambda _: self._cari_barang())
        tk.Label(row1, text="Kode/Nama", bg=C_PUTIH, fg="#888",
                 font=("Times New Roman", 8)).pack(side="left")

        tk.Button(row1, text="Cari", command=self._cari_barang,
                  bg=C_HEADER, fg=C_KUNING, font=FONT_B,
                  relief="flat", padx=12, cursor="hand2").pack(side="right")

        # Kategori tabs
        row2 = tk.Frame(parent, bg=C_PUTIH)
        row2.pack(fill="x", pady=(0, 6))
        kats = [("semua","Semua"),("makanan","🍱Makanan"),("minuman","🥤Minuman"),
                ("snack","🍿Snack"),("sachet","☕Sachet"),("sabun","🧺Cuci"),
                ("kosmetik","💄Kosmetik"),("lainnya","📦Lainnya")]
        self._kat_btns = {}
        for kat, label in kats:
            btn = tk.Button(row2, text=label, font=("Times New Roman", 8),
                            relief="flat", cursor="hand2", padx=6, pady=2,
                            command=lambda k=kat: self._filter_kat(k))
            btn.pack(side="left", padx=2)
            self._kat_btns[kat] = btn
        self._update_kat_btn("semua")

        # List barang
        cols = ("kode", "nama", "harga", "kat")
        self.tree_barang = ttk.Treeview(parent, columns=cols,
                                        show="headings", height=8)
        for col, h, w in [("kode","Kode",60),("nama","Nama Barang",300),
                           ("harga","Harga",100),("kat","Kategori",90)]:
            self.tree_barang.heading(col, text=h)
            self.tree_barang.column(col, width=w, anchor="w" if col=="nama" else "center")
        sb = ttk.Scrollbar(parent, orient="vertical",
                           command=self.tree_barang.yview)
        self.tree_barang.configure(yscrollcommand=sb.set)
        self.tree_barang.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.tree_barang.bind("<Double-1>", self._on_barang_dbl)
        self.tree_barang.bind("<Return>",   self._on_barang_dbl)

        style = ttk.Style()
        style.configure("Treeview", font=FONT_J, rowheight=22)
        style.configure("Treeview.Heading", font=FONT_B,
                        background=C_MID, foreground=C_PUTIH)
        style.map("Treeview", background=[("selected", C_LIGHT)])

    def _build_cart(self, parent):
        cols = ("nama", "qty", "satuan", "subtotal")
        self.tree_cart = ttk.Treeview(parent, columns=cols,
                                      show="headings", height=12)
        for col, h, w, anc in [
            ("nama","Nama Barang",260,"w"),("qty","Qty",50,"center"),
            ("satuan","Harga/pcs",110,"e"),("subtotal","Subtotal",120,"e")]:
            self.tree_cart.heading(col, text=h)
            self.tree_cart.column(col, width=w, anchor=anc)
        sb2 = ttk.Scrollbar(parent, orient="vertical",
                             command=self.tree_cart.yview)
        self.tree_cart.configure(yscrollcommand=sb2.set)
        self.tree_cart.pack(side="left", fill="both", expand=True)
        sb2.pack(side="right", fill="y")
        self.tree_cart.bind("<Double-1>", self._on_cart_edit)

        btn_row = tk.Frame(parent, bg=C_PUTIH)
        btn_row.pack(fill="x", pady=(6, 0))
        tk.Button(btn_row, text="+ Tambah 1", command=self._tambah1,
                  bg=C_MID, fg="white", font=FONT_J,
                  relief="flat", cursor="hand2", padx=8).pack(side="left", padx=2)
        tk.Button(btn_row, text="− Kurangi 1", command=self._kurangi1,
                  bg="#1976D2", fg="white", font=FONT_J,
                  relief="flat", cursor="hand2", padx=8).pack(side="left", padx=2)
        tk.Button(btn_row, text="🗑 Hapus", command=self._hapus_item,
                  bg="#EF5350", fg="white", font=FONT_J,
                  relief="flat", cursor="hand2", padx=8).pack(side="left", padx=2)

    def _build_action(self, parent):
        inner = tk.Frame(parent, bg=C_LIGHT, pady=8)
        inner.pack(fill="x", padx=10)

        tk.Label(inner, text="Total Belanja:", bg=C_LIGHT, fg="#555",
                 font=FONT_J).pack(side="left", padx=8)
        self.lbl_total = tk.Label(inner, text="Rp 0", bg=C_LIGHT, fg=C_HEADER,
                                   font=("Times New Roman", 17, "bold"))
        self.lbl_total.pack(side="left", padx=4)

        tk.Button(inner, text="🗑 Kosongkan", command=self._kosongkan,
                  bg="#EF5350", fg="white", font=FONT_J,
                  relief="flat", cursor="hand2", padx=10).pack(side="right", padx=4)
        tk.Button(inner, text="➕ Next Item", command=self._next_item,
                  bg=C_MID, fg="white", font=FONT_B,
                  relief="flat", cursor="hand2", padx=14).pack(side="right", padx=4)
        tk.Button(inner, text="💳  BAYAR", command=self._buka_bayar,
                  bg=C_KUNING, fg=C_HEADER, font=("Times New Roman", 12, "bold"),
                  relief="flat", cursor="hand2", padx=18).pack(side="right", padx=4)

    # ── RIGHT PANEL ───────────────────────────────────────────────────────────
    def _build_right(self, parent):
        right = tk.Frame(parent, bg=C_BG)
        right.grid(row=0, column=1, sticky="nsew")
        right.rowconfigure(1, weight=1)
        right.rowconfigure(2, weight=2)
        right.columnconfigure(0, weight=1)

        # Struk
        f_struk = self._frame_box(right, "🧾  Preview Struk", 0, 0, rowspan=2)
        self.txt_struk = tk.Text(f_struk, width=40, height=22,
                                  font=("Courier New", 8),
                                  bg="#FAFAFA", fg="#222", relief="flat",
                                  state="disabled", wrap="none")
        sb_s = ttk.Scrollbar(f_struk, orient="vertical",
                              command=self.txt_struk.yview)
        self.txt_struk.configure(yscrollcommand=sb_s.set)
        self.txt_struk.pack(side="left", fill="both", expand=True)
        sb_s.pack(side="right", fill="y")

        # Tombol bawah struk
        f_btn = self._frame_box(right, "", 2, 0, pady_inner=4)
        tk.Button(f_btn, text="🖨  Cetak Struk (Jendela Baru)",
                  command=self._cetak_struk,
                  bg=C_MID, fg="white", font=FONT_B,
                  relief="flat", cursor="hand2").pack(fill="x", pady=3)
        tk.Button(f_btn, text="📊  Tampilkan Grafik Penjualan",
                  command=self._buka_grafik,
                  bg=C_HEADER, fg=C_KUNING, font=FONT_B,
                  relief="flat", cursor="hand2").pack(fill="x", pady=3)
        tk.Button(f_btn, text="📁  Buka Folder File Excel",
                  command=self._buka_folder,
                  bg="#37474F", fg="white", font=FONT_J,
                  relief="flat", cursor="hand2").pack(fill="x", pady=3)

    # ── LOGIKA ────────────────────────────────────────────────────────────────
    def _set_tanggal_hari_ini(self):
        self.var_tgl.set(datetime.today().strftime("%Y-%m-%d"))

    def _muat_barang(self):
        try:
            self._barang_list = FileHandler.baca_barang()
            self._filtered    = list(self._barang_list)
            self._render_barang()
            self._set_status(f"✅ {len(self._barang_list)} barang dimuat dari items_barang.xlsx")
        except FileNotFoundError as e:
            messagebox.showerror("File Tidak Ditemukan", str(e))
        except Exception as e:
            messagebox.showerror("Error", f"Gagal memuat barang:\n{e}")

    def _render_barang(self):
        self.tree_barang.delete(*self.tree_barang.get_children())
        for b in self._filtered:
            self.tree_barang.insert("", "end", iid=b.kode,
                                    values=(b.kode, b.nama,
                                            f"Rp{b.harga:,}".replace(",","."),
                                            b.kategori))

    def _cari_barang(self):
        q = self.var_kode.get().strip().lower()
        if not q:
            self._filtered = list(self._barang_list)
        else:
            self._filtered = [b for b in self._barang_list
                               if q in b.kode.lower() or q in b.nama.lower()]
        if self._kat_aktif != "semua":
            self._filtered = [b for b in self._filtered
                               if b.kategori == self._kat_aktif]
        self._render_barang()
        if len(self._filtered) == 1:
            self._tambah_ke_cart(self._filtered[0])
            self.var_kode.set("")

    def _filter_kat(self, kat: str):
        self._kat_aktif = kat
        self._update_kat_btn(kat)
        self.var_kode.set("")
        self._filtered = [b for b in self._barang_list
                          if kat == "semua" or b.kategori == kat]
        self._render_barang()

    def _update_kat_btn(self, aktif: str):
        for k, btn in self._kat_btns.items():
            if k == aktif:
                btn.configure(bg=C_HEADER, fg=C_KUNING)
            else:
                btn.configure(bg=C_LIGHT, fg=C_TEXT)

    def _on_barang_dbl(self, _=None):
        sel = self.tree_barang.focus()
        if not sel:
            return
        barang = next((b for b in self._barang_list if b.kode == sel), None)
        if barang:
            self._tambah_ke_cart(barang)
            self.var_kode.set("")

    def _tambah_ke_cart(self, barang: Barang):
        tgl = self.var_tgl.get()
        if not self._transaksi:
            self._transaksi = Transaksi(tgl, "cash")
        self._transaksi.tambah_item(barang, 1)
        self._render_cart()

    def _render_cart(self):
        self.tree_cart.delete(*self.tree_cart.get_children())
        if not self._transaksi:
            self.lbl_total.configure(text="Rp 0")
            return
        for it in self._transaksi.items:
            self.tree_cart.insert("", "end", iid=it.barang.kode,
                                  values=(it.barang.nama, it.qty,
                                          f"Rp{it.barang.harga:,}".replace(",","."),
                                          f"Rp{it.subtotal:,}".replace(",",".")))
        total = self._transaksi.total
        self.lbl_total.configure(text=f"Rp{total:,}".replace(",","."))

    def _on_cart_edit(self, _=None):
        pass  # Double-click di cart → bisa extend untuk edit qty langsung

    def _tambah1(self):
        sel = self.tree_cart.focus()
        if not sel or not self._transaksi:
            return
        barang = next((b for b in self._barang_list if b.kode == sel), None)
        if barang:
            self._transaksi.tambah_item(barang, 1)
            self._render_cart()

    def _kurangi1(self):
        sel = self.tree_cart.focus()
        if not sel or not self._transaksi:
            return
        for it in self._transaksi.items:
            if it.barang.kode == sel:
                it.qty -= 1
                if it.qty <= 0:
                    self._transaksi.hapus_item(sel)
                break
        self._render_cart()

    def _hapus_item(self):
        sel = self.tree_cart.focus()
        if sel and self._transaksi:
            self._transaksi.hapus_item(sel)
            self._render_cart()

    def _kosongkan(self):
        if self._transaksi and self._transaksi.items:
            if messagebox.askyesno("Konfirmasi", "Kosongkan keranjang belanja?"):
                self._transaksi = None
                self._render_cart()

    def _next_item(self):
        self.var_kode.set("")
        self._filter_kat("semua")
        self.tree_barang.focus_set()

    # ── PEMBAYARAN ────────────────────────────────────────────────────────────
    def _buka_bayar(self):
        if not self._transaksi or not self._transaksi.items:
            messagebox.showwarning("Keranjang Kosong", "Tambahkan barang terlebih dahulu!")
            return

        win = tk.Toplevel(self)
        win.title("Pembayaran")
        win.geometry("420x420")
        win.configure(bg=C_BG)
        win.grab_set()
        win.resizable(False, False)

        tk.Label(win, text="Pilih Metode Pembayaran", bg=C_HEADER, fg=C_KUNING,
                 font=FONT_BIG).pack(fill="x", ipady=10)
        tk.Label(win, text="Total yang harus dibayar:", bg=C_BG, fg="#555",
                 font=FONT_J).pack(pady=(14, 2))
        tk.Label(win, text=f"Rp{self._transaksi.total:,}".replace(",","."),
                 bg=C_BG, fg=C_HEADER,
                 font=("Times New Roman", 24, "bold")).pack()

        var_metode = tk.StringVar(value="")

        frm = tk.Frame(win, bg=C_BG)
        frm.pack(pady=12)
        for metode, icon in [("qris","📱  QRIS"),("cash","💵  Cash")]:
            tk.Radiobutton(frm, text=icon, variable=var_metode, value=metode,
                           bg=C_BG, fg=C_TEXT, font=FONT_BIG,
                           activebackground=C_LIGHT,
                           selectcolor=C_LIGHT).pack(side="left", padx=18, ipadx=10, ipady=6)

        # Cash input
        frm_cash = tk.Frame(win, bg=C_BG)
        frm_cash.pack(pady=4)
        tk.Label(frm_cash, text="Uang yang dibayar:", bg=C_BG,
                 font=FONT_J).pack()
        var_cash = tk.StringVar()
        lbl_kembali = tk.Label(frm_cash, text="", bg=C_BG, fg=C_MID, font=FONT_B)
        e_cash = tk.Entry(frm_cash, textvariable=var_cash, width=20, font=FONT_J,
                          bg="#F0F8FF", relief="solid")
        e_cash.pack(pady=4)
        lbl_kembali.pack()

        def hitung_kembali(*_):
            try:
                bayar = int(var_cash.get())
                kembali = bayar - self._transaksi.total
                if kembali < 0:
                    lbl_kembali.configure(text=f"⚠ Kurang: Rp{abs(kembali):,}".replace(",","."), fg="red")
                else:
                    lbl_kembali.configure(text=f"✅ Kembalian: Rp{kembali:,}".replace(",","."), fg="#2E7D32")
            except ValueError:
                lbl_kembali.configure(text="")

        def on_metode(*_):
            if var_metode.get() == "cash":
                frm_cash.pack(pady=4)
            else:
                frm_cash.pack_forget()

        var_metode.trace_add("write", on_metode)
        var_cash.trace_add("write", hitung_kembali)
        frm_cash.pack_forget()

        def selesai():
            metode = var_metode.get()
            if not metode:
                messagebox.showwarning("Pilih Metode", "Pilih metode pembayaran!", parent=win)
                return
            uang = 0
            if metode == "cash":
                try:
                    uang = int(var_cash.get())
                except ValueError:
                    messagebox.showerror("Error", "Masukkan jumlah uang!", parent=win)
                    return
                if uang < self._transaksi.total:
                    messagebox.showerror("Error", "Uang kurang dari total belanja!", parent=win)
                    return
            self._transaksi.metode = metode
            self._uang_bayar = uang
            self._selesai_bayar()
            win.destroy()

        tk.Button(win, text="✅  Selesai & Buat Struk", command=selesai,
                  bg=C_HEADER, fg=C_KUNING, font=FONT_BIG,
                  relief="flat", cursor="hand2", pady=10).pack(fill="x", padx=24, pady=14)
        tk.Button(win, text="Batal", command=win.destroy,
                  bg="white", fg="#888", font=FONT_J,
                  relief="flat", cursor="hand2").pack()

    def _selesai_bayar(self):
        try:
            path = FileHandler.simpan_transaksi(self._transaksi)
            struk = self._transaksi.buat_struk(self._uang_bayar)
            self._tampil_struk(struk)
            self._trx_count += 1
            self._trx_total += self._transaksi.total
            self.lbl_trx_count.configure(
                text=f"Transaksi: {self._trx_count}  |  Total: Rp{self._trx_total:,}".replace(",","."))
            messagebox.showinfo("Berhasil",
                                f"✅ Pembayaran berhasil!\nDisimpan ke: {os.path.basename(path)}")
            self._transaksi = None
            self._render_cart()
        except Exception as e:
            messagebox.showerror("Error Simpan", str(e))

    def _tampil_struk(self, struk: str):
        self.txt_struk.configure(state="normal")
        self.txt_struk.delete("1.0", "end")
        self.txt_struk.insert("end", struk)
        self.txt_struk.configure(state="disabled")

    def _cetak_struk(self):
        struk = self.txt_struk.get("1.0", "end").strip()
        if not struk:
            messagebox.showinfo("Info", "Belum ada struk untuk dicetak.")
            return
        win = tk.Toplevel(self)
        win.title("Struk Pembelian")
        win.geometry("340x600")
        win.configure(bg="white")
        txt = tk.Text(win, font=("Courier New", 9), bg="white",
                      relief="flat", padx=10, pady=10)
        txt.pack(fill="both", expand=True)
        txt.insert("end", struk)
        txt.configure(state="disabled")
        tk.Button(win, text="🖨 Print", command=lambda: win.event_generate("<<Print>>"),
                  bg=C_HEADER, fg=C_KUNING, font=FONT_B,
                  relief="flat").pack(fill="x", pady=6, padx=20)

    def _buka_grafik(self):
        win = tk.Toplevel(self)
        win.title("Grafik Penjualan")
        win.geometry("800x420")
        win.configure(bg=C_BG)
        tk.Label(win, text="📊 Grafik Penjualan Bulanan", bg=C_HEADER, fg=C_KUNING,
                 font=FONT_BIG).pack(fill="x", ipady=8)
        frm = tk.Frame(win, bg=C_BG)
        frm.pack(fill="both", expand=True, padx=10, pady=10)
        GrafikPenjualan.tampilkan(frm, self.var_tgl.get())

    def _buka_folder(self):
        folder = FileHandler.BASE_DIR
        try:
            if sys.platform == "win32":
                os.startfile(folder)
            elif sys.platform == "darwin":
                os.system(f'open "{folder}"')
            else:
                os.system(f'xdg-open "{folder}"')
        except Exception as e:
            messagebox.showinfo("Lokasi File", folder)

    def _set_status(self, msg: str):
        # Tampilkan di title bar
        self.title(f"Jack(owi) Mart — {msg}")


# ═════════════════════════════════════════════════════════════════════════════
#  ENTRY POINT
# ═════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    app = KasirApp()
    app.mainloop()


