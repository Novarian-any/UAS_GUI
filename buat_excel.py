import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
import random
import os

# ─── WARNA ───────────────────────────────────────────────────────────────────
BIRU_TUA   = "0D47A1"
BIRU_MID   = "1565C0"
BIRU_MUDA  = "BBDEFB"
KUNING     = "FFD54F"
PUTIH      = "FFFFFF"
ABU        = "F5F5F5"
ABU_BORDER = "BDBDBD"

def border_tipis():
    s = Side(style='thin', color=ABU_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)

# ─── DATA 100 BARANG NYATA ────────────────────────────────────────────────────
BARANG = [
    # MAKANAN (001-020)
    ("001","Indomie Goreng Original 85gr",          3200,  "makanan"),
    ("002","Indomie Soto Ayam 70gr",                3200,  "makanan"),
    ("003","Indomie Rendang 85gr",                  3500,  "makanan"),
    ("004","Mie Sedaap Goreng 87gr",                3000,  "makanan"),
    ("005","Sarimi Soto Koya 75gr",                 2500,  "makanan"),
    ("006","Pop Mie Ayam Bawang 75gr",              4500,  "makanan"),
    ("007","Beras Ramos 1kg",                      13000,  "makanan"),
    ("008","Gula Pasir Gulaku 1kg",                16000,  "makanan"),
    ("009","Minyak Goreng Bimoli 1L",              21000,  "makanan"),
    ("010","Kecap Manis Bango 135ml",               9000,  "makanan"),
    ("011","Kecap ABC Manis 275ml",                13000,  "makanan"),
    ("012","Saus Sambal ABC 135ml",                 9000,  "makanan"),
    ("013","Royco Bumbu Penyedap Ayam 250gr",      12000,  "makanan"),
    ("014","Masako Penyedap Sapi 100gr",            7500,  "makanan"),
    ("015","Tepung Terigu Segitiga Biru 1kg",      13000,  "makanan"),
    ("016","Saos Tomat Del Monte 340gr",           15000,  "makanan"),
    ("017","Madu TJ Murni 350gr",                  35000,  "makanan"),
    ("018","Roti Tawar Sari Roti Original",        16000,  "makanan"),
    ("019","Biskuat Cokelat 100gr",                 8500,  "makanan"),
    ("020","SilverQueen Chunky Bar 62gr",           9500,  "makanan"),

    # MINUMAN (021-040)
    ("021","Aqua Botol 600ml",                      4500,  "minuman"),
    ("022","Aqua Gelas 240ml",                      1500,  "minuman"),
    ("023","Teh Botol Sosro 450ml",                 6000,  "minuman"),
    ("024","Frestea Green 500ml",                   7000,  "minuman"),
    ("025","Coca-Cola Kaleng 330ml",                8500,  "minuman"),
    ("026","Sprite Kaleng 330ml",                   8500,  "minuman"),
    ("027","Fanta Strawberry Kaleng 330ml",         8500,  "minuman"),
    ("028","Pocari Sweat 500ml",                   10000,  "minuman"),
    ("029","Mizone Jeruk Lemon 500ml",              8000,  "minuman"),
    ("030","Nu Green Tea Original 330ml",           6500,  "minuman"),
    ("031","ABC Juice Guava 250ml",                 6000,  "minuman"),
    ("032","Minute Maid Pulpy Orange 350ml",        7500,  "minuman"),
    ("033","Bear Brand Susu Steril 189ml",         10000,  "minuman"),
    ("034","Ultra Milk Full Cream 250ml",            7000,  "minuman"),
    ("035","Yakult 65ml",                           4000,  "minuman"),
    ("036","Adem Sari Ching Ku 300ml",              7500,  "minuman"),
    ("037","Good Day Cappuccino 250ml",             7000,  "minuman"),
    ("038","Es Teh Indonesia 350ml",                6000,  "minuman"),
    ("039","Floridina Orange 350ml",                5500,  "minuman"),
    ("040","Pepsi Cola Kaleng 330ml",               8000,  "minuman"),

    # SNACK (041-055)
    ("041","Chitato Sapi Panggang 68gr",            9500,  "snack"),
    ("042","Chitato Ayam Lada Hitam 68gr",          9500,  "snack"),
    ("043","Lay's Rumput Laut 68gr",               10000,  "snack"),
    ("044","Pringles Original 107gr",              25000,  "snack"),
    ("045","Oreo Original 137gr",                  12000,  "snack"),
    ("046","Roma Marie Susu 225gr",                12000,  "snack"),
    ("047","Khong Guan Assorted 300gr",            27000,  "snack"),
    ("048","Wafer Tango Cokelat 150gr",            13000,  "snack"),
    ("049","Superstar Snack Jagung 60gr",            5000,  "snack"),
    ("050","Qtela Singkong Original 65gr",           8000,  "snack"),
    ("051","Makaroni Ngehe 55gr",                   5500,  "snack"),
    ("052","Kacang Garuda Original 100gr",           9000,  "snack"),
    ("053","Sukro Kacang Atom 120gr",               8500,  "snack"),
    ("054","Cheetos Twist Keju 55gr",               7500,  "snack"),
    ("055","Lays BBQ 68gr",                        10000,  "snack"),

    # MINUMAN SACHET (056-065)
    ("056","Kopi Kapal Api Special Mix 25gr",       2500,  "sachet"),
    ("057","Nescafe Classic 2gr (1 sachet)",        1500,  "sachet"),
    ("058","Nescafe 3in1 Original 20gr",            3500,  "sachet"),
    ("059","Good Day 3in1 Cappuccino 20gr",         3000,  "sachet"),
    ("060","Milo Activ-Go Sachet 22gr",             3500,  "sachet"),
    ("061","Energen Sereal Cokelat 38gr",            4000,  "sachet"),
    ("062","Sariwangi Teh Celup 1,6gr (1 sachet)",   500,  "sachet"),
    ("063","Teh Pucuk 2gr (1 sachet)",               500,  "sachet"),
    ("064","Susu Dancow Full Cream Sachet 26gr",    4000,  "sachet"),
    ("065","Ovomaltine Sachet 18gr",                5000,  "sachet"),

    # PRODUK CUCI PAKAIAN (066-075)
    ("066","Rinso Anti Noda 1kg",                  22000,  "sabun"),
    ("067","Daia Deterjen Bunga 800gr",            14000,  "sabun"),
    ("068","Attack Deterjen Easy Matic 1kg",       22000,  "sabun"),
    ("069","Surf Deterjen Parfum Bunga 800gr",     15000,  "sabun"),
    ("070","Soklin Liquid Detergent 750ml",        23000,  "sabun"),
    ("071","Molto Ultra Sekali Bilas 800ml",       22000,  "sabun"),
    ("072","Downy Pelembut Pakaian Fresh 900ml",   23000,  "sabun"),
    ("073","Bayclin Pemutih Pakaian 800ml",        16000,  "sabun"),
    ("074","Wipol Karbol Kamper 780ml",            18000,  "sabun"),
    ("075","Sunlight Pencuci Piring Jeruk 755ml",  17000,  "sabun"),

    # KOSMETIK & PERAWATAN (076-090)
    ("076","Wardah Lightening BB Cream 15gr",      35000,  "kosmetik"),
    ("077","Emina Creamatte Lipstick 4gr",         32000,  "kosmetik"),
    ("078","Pond's White Beauty Day Cream 50gr",   43000,  "kosmetik"),
    ("079","Nivea Body Lotion Extra White 200ml",  38000,  "kosmetik"),
    ("080","Citra Body Lotion Sakura 230ml",       29000,  "kosmetik"),
    ("081","Vaseline Intensive Care Aloe 200ml",   35000,  "kosmetik"),
    ("082","Dove Body Wash Deeply Nourishing 450ml",42000, "kosmetik"),
    ("083","Lifebuoy Sabun Cair Total 10 250ml",   22000,  "kosmetik"),
    ("084","Nuvo Sabun Cair Family 250ml",         19000,  "kosmetik"),
    ("085","Pantene Shampoo 2in1 70ml",            12000,  "kosmetik"),
    ("086","Clear Men Shampoo Ice Cool 70ml",      12000,  "kosmetik"),
    ("087","Garnier Men Acno Fight 100ml",         38000,  "kosmetik"),
    ("088","Hazeline Snow Original 60gr",          25000,  "kosmetik"),
    ("089","Gatsby Pomade Water Gloss 75gr",       28000,  "kosmetik"),
    ("090","Pepsodent Action 123 Pasta Gigi 190gr",17000,  "kosmetik"),

    # LAINNYA (091-100)
    ("091","Baterai ABC AA 2pcs",                   9000,  "lainnya"),
    ("092","Baterai ABC AAA 2pcs",                  8500,  "lainnya"),
    ("093","Tisu Paseo 250 Lembar",                15000,  "lainnya"),
    ("094","Tisu Basah Mitu 50 Lembar",            12000,  "lainnya"),
    ("095","Pembalut Wanita Charm 16pcs",          22000,  "lainnya"),
    ("096","Pembalut Laurier Active 16pcs",        23000,  "lainnya"),
    ("097","Korek Api RI 1 Box",                    2000,  "lainnya"),
    ("098","Kantong Plastik Kresek Merah (50pcs)",  8000,  "lainnya"),
    ("099","Pulpen BallPoint Pilot 1 pcs",          5000,  "lainnya"),
    ("100","Buku Tulis Kiky 38 Lembar",             5500,  "lainnya"),
]

# ─── BUAT FILE ITEMS ──────────────────────────────────────────────────────────
def buat_file_items(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daftar Barang"

    # Header baris toko
    ws.merge_cells("A1:D1")
    ws["A1"] = "JACK(OWI) MART - DAFTAR BARANG"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True, color=KUNING)
    ws["A1"].fill = PatternFill("solid", fgColor=BIRU_TUA)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Header kolom
    headers = ["Kode", "Nama Barang", "Harga (Rp)", "Kategori"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=col, value=h)
        c.font = Font(name="Times New Roman", size=11, bold=True, color=PUTIH)
        c.fill = PatternFill("solid", fgColor=BIRU_MID)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_tipis()
    ws.row_dimensions[2].height = 22

    # Data barang
    for i, (kode, nama, harga, kat) in enumerate(BARANG):
        row = i + 3
        bg = ABU if i % 2 == 0 else PUTIH
        vals = [kode, nama, harga, kat]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.font = Font(name="Times New Roman", size=10)
            c.fill = PatternFill("solid", fgColor=bg)
            c.border = border_tipis()
            if col == 1:
                c.alignment = Alignment(horizontal="center")
            elif col == 3:
                c.alignment = Alignment(horizontal="right")
                c.number_format = '#,##0'
            else:
                c.alignment = Alignment(horizontal="left")

    # Total row
    total_row = len(BARANG) + 3
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(name="Times New Roman", bold=True)
    ws.cell(row=total_row, column=1).fill = PatternFill("solid", fgColor=BIRU_MUDA)
    ws.cell(row=total_row, column=2, value=f"{len(BARANG)} Jenis Barang")
    ws.cell(row=total_row, column=2).font = Font(name="Times New Roman", bold=True)
    ws.cell(row=total_row, column=2).fill = PatternFill("solid", fgColor=BIRU_MUDA)
    ws.cell(row=total_row, column=3, value=f"=SUM(C3:C{total_row-1})")
    ws.cell(row=total_row, column=3).font = Font(name="Times New Roman", bold=True)
    ws.cell(row=total_row, column=3).fill = PatternFill("solid", fgColor=BIRU_MUDA)
    ws.cell(row=total_row, column=3).number_format = '#,##0'

    # Lebar kolom
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14

    # Freeze header
    ws.freeze_panes = "A3"

    wb.save(path)
    print(f"[OK] File items dibuat: {path}")

# ─── BUAT FILE PENJUALAN TEMPLATE ─────────────────────────────────────────────
def buat_file_penjualan(path):
    wb = openpyxl.Workbook()

    # ── Sheet 1: Log Transaksi ──────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Log Transaksi"

    ws1.merge_cells("A1:H1")
    ws1["A1"] = "JACK(OWI) MART - LOG TRANSAKSI PENJUALAN"
    ws1["A1"].font = Font(name="Times New Roman", size=14, bold=True, color=KUNING)
    ws1["A1"].fill = PatternFill("solid", fgColor=BIRU_TUA)
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 30

    h2 = ["No","Tanggal","No Transaksi","Kode Barang","Nama Barang","Qty","Harga Satuan","Subtotal"]
    for col, h in enumerate(h2, 1):
        c = ws1.cell(row=2, column=col, value=h)
        c.font = Font(name="Times New Roman", size=10, bold=True, color=PUTIH)
        c.fill = PatternFill("solid", fgColor=BIRU_MID)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border_tipis()
    ws1.row_dimensions[2].height = 28

    # Contoh data dummy 3 hari
    contoh = []
    base_date = datetime.today() - timedelta(days=2)
    no = 1
    for d in range(3):
        tgl = (base_date + timedelta(days=d)).strftime("%Y-%m-%d")
        n_trx = random.randint(3, 5)
        for t in range(n_trx):
            no_trx = f"TRX/{(base_date+timedelta(days=d)).strftime('%Y%m%d')}/{str(t+1).zfill(3)}"
            n_item = random.randint(1, 4)
            pilihan = random.sample(BARANG, n_item)
            for kode, nama, harga, _ in pilihan:
                qty = random.randint(1, 5)
                contoh.append((no, tgl, no_trx, kode, nama, qty, harga))
                no += 1

    for i, (no_r, tgl, no_trx, kode, nama, qty, harga) in enumerate(contoh):
        row = i + 3
        bg = ABU if i % 2 == 0 else PUTIH
        subtotal_formula = f"=G{row}*F{row}"
        vals = [no_r, tgl, no_trx, kode, nama, qty, harga, subtotal_formula]
        for col, val in enumerate(vals, 1):
            c = ws1.cell(row=row, column=col, value=val)
            c.font = Font(name="Times New Roman", size=10)
            c.fill = PatternFill("solid", fgColor=bg)
            c.border = border_tipis()
            if col in [1, 4, 6]:
                c.alignment = Alignment(horizontal="center")
            elif col in [7, 8]:
                c.alignment = Alignment(horizontal="right")
                c.number_format = '#,##0'
            elif col == 2:
                c.alignment = Alignment(horizontal="center")
                c.number_format = 'YYYY-MM-DD'

    last_data = len(contoh) + 2
    total_r = last_data + 1
    ws1.cell(row=total_r, column=5, value="GRAND TOTAL")
    ws1.cell(row=total_r, column=5).font = Font(name="Times New Roman", bold=True)
    ws1.cell(row=total_r, column=5).fill = PatternFill("solid", fgColor=BIRU_MUDA)
    ws1.cell(row=total_r, column=8, value=f"=SUM(H3:H{last_data})")
    ws1.cell(row=total_r, column=8).font = Font(name="Times New Roman", bold=True)
    ws1.cell(row=total_r, column=8).fill = PatternFill("solid", fgColor=BIRU_MUDA)
    ws1.cell(row=total_r, column=8).number_format = '#,##0'

    lebar1 = [6, 14, 26, 10, 38, 6, 15, 15]
    for i, w in enumerate(lebar1, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = "A3"

    # ── Sheet 2: Rekap Harian ───────────────────────────────────────────────
    ws2 = wb.create_sheet("Rekap Harian")

    ws2.merge_cells("A1:E1")
    ws2["A1"] = "REKAP PENJUALAN HARIAN"
    ws2["A1"].font = Font(name="Times New Roman", size=13, bold=True, color=KUNING)
    ws2["A1"].fill = PatternFill("solid", fgColor=BIRU_TUA)
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    h3 = ["Tanggal","Jumlah Transaksi","Total Item Terjual","Total Pendapatan (Rp)","Catatan"]
    for col, h in enumerate(h3, 1):
        c = ws2.cell(row=2, column=col, value=h)
        c.font = Font(name="Times New Roman", size=10, bold=True, color=PUTIH)
        c.fill = PatternFill("solid", fgColor=BIRU_MID)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border_tipis()
    ws2.row_dimensions[2].height = 28

    # Isi rekap dari contoh data
    from collections import defaultdict
    rekap = defaultdict(lambda: {"trx": set(), "item": 0, "total": 0})
    for _, tgl, no_trx, _, _, qty, harga in contoh:
        rekap[tgl]["trx"].add(no_trx)
        rekap[tgl]["item"] += qty
        rekap[tgl]["total"] += qty * harga

    for i, (tgl, info) in enumerate(sorted(rekap.items())):
        row = i + 3
        bg = ABU if i % 2 == 0 else PUTIH
        vals = [tgl, len(info["trx"]), info["item"], info["total"], ""]
        for col, val in enumerate(vals, 1):
            c = ws2.cell(row=row, column=col, value=val)
            c.font = Font(name="Times New Roman", size=10)
            c.fill = PatternFill("solid", fgColor=bg)
            c.border = border_tipis()
            if col == 1:
                c.alignment = Alignment(horizontal="center")
                c.number_format = 'YYYY-MM-DD'
            elif col == 4:
                c.alignment = Alignment(horizontal="right")
                c.number_format = '#,##0'
            else:
                c.alignment = Alignment(horizontal="center")

    last_rekap = len(rekap) + 2
    for col in range(1, 6):
        c = ws2.cell(row=last_rekap+1, column=col)
        c.fill = PatternFill("solid", fgColor=BIRU_MUDA)
        c.border = border_tipis()
    ws2.cell(row=last_rekap+1, column=1, value="TOTAL")
    ws2.cell(row=last_rekap+1, column=1).font = Font(name="Times New Roman", bold=True)
    ws2.cell(row=last_rekap+1, column=4, value=f"=SUM(D3:D{last_rekap})")
    ws2.cell(row=last_rekap+1, column=4).font = Font(name="Times New Roman", bold=True)
    ws2.cell(row=last_rekap+1, column=4).number_format = '#,##0'

    lebar2 = [16, 20, 20, 24, 20]
    for i, w in enumerate(lebar2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A3"

    wb.save(path)
    print(f"[OK] File penjualan dibuat: {path}")

if __name__ == "__main__":
    out_dir = os.path.dirname(os.path.abspath(__file__))
    buat_file_items(os.path.join(out_dir, "items_barang.xlsx"))
    buat_file_penjualan(os.path.join(out_dir, "penjualan.xlsx"))
    print("\nSemua file Excel berhasil dibuat!")
